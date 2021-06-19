from openpyxl.workbook import Workbook
from openpyxl import load_workbook
from openpyxl.chart import PieChart, PieChart3D, Reference, BarChart, BarChart3D, LineChart, LineChart3D

wb = load_workbook('hello1.xlsx')
ws = wb.active

# Determine the type of chart you want
# chart = PieChart()
# simply chanage to PieChart3D and keep the following the same
# chart = PieChart3D()
# chart = BarChart()
# chart = BarChart3D()
chart = LineChart()

# designate labels and data
labels = Reference(ws, min_col=1, max_col=1, min_row=2, max_row=10) # remove the header
data = Reference(ws, min_col=3, min_row=1, max_row=10) # keep the header

# put it all together
'''
Args:
titles_from_data: grab the title from the data header
'''
chart.add_data(data, titles_from_data=True)
chart.set_categories(labels)

# add a title
chart.title = 'Employee Salaries'

# place the chart anchoring to top left cell
ws.add_chart(chart, 'E2')

# save
wb.save('hello2.xlsx')
print('File saved.')